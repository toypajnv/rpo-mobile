from __future__ import annotations

from datetime import datetime, timezone
import json
import tempfile
import unittest

from openpyxl import load_workbook

from app.models import PermitRecord
from app.services.exporter import MAX_HEADERS, build_export


class MaxCompatibleExportTests(unittest.TestCase):
    def _record(self) -> PermitRecord:
        fields = {
            "AT": {"field_value": "29.09.2026 10:45", "comment": ""},
            "AU": {"field_value": "29.09.2026 10:49", "comment": ""},
            "AV": {"field_value": "29.09.2026 10:49", "comment": ""},
            "AX": {"field_value": "29.09.2026 10:50", "comment": ""},
            "AY": {"field_value": "29.09.2026 10:51", "comment": ""},
            "AZ": {"field_value": "29.09.2026 11:00", "comment": "Нет допуска к месту работ"},
            "BA": {"field_value": "29.09.2026 11:10", "comment": ""},
            "BC": {"field_value": "29.09.2026 11:20", "comment": ""},
            "BD": {"field_value": "29.09.2026 11:25", "comment": ""},
            "BE": {"field_value": "30.09.2026", "comment": "Работы выполнены не в полном объеме"},
            "BF": {"field_value": "29.09.2026 11:30", "comment": ""},
            "BG": {"field_value": "29.09.2026 11:35", "comment": ""},
            "BH": {"field_value": "29.09.2026 11:40", "comment": ""},
            "RI": {"field_value": "29.09.2026 11:05", "comment": "Замена исполнителя"},
        }
        return PermitRecord(
            id=199,
            permit_number="СН-041285",
            worker_name="Костин В.А.",
            structural_unit="ЦДПН-1",
            device_id="device-1",
            data_json=json.dumps(fields, ensure_ascii=False),
            updated_at=datetime(2026, 9, 30, 4, 35, 31, tzinfo=timezone.utc),
        )

    def test_xlsx_matches_max_upload_layout(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            xlsx_path, json_path = build_export([self._record()], tmp, 29)

            self.assertTrue(xlsx_path.name.startswith("Шаблон_загрузки_данных_MAX_заполненный_"))
            self.assertTrue(json_path.exists())

            wb = load_workbook(xlsx_path)
            self.assertEqual(wb.sheetnames, ["Выгрузка РПО", "Инструкция"])
            ws = wb["Выгрузка РПО"]

            self.assertEqual([ws.cell(1, col).value for col in range(1, 22)], MAX_HEADERS)
            self.assertEqual(ws.tables["MaxUploadTable"].ref, "A1:U501")
            self.assertEqual(len(ws.data_validations.dataValidation), 6)

            self.assertEqual(ws["A2"].value, "новый")
            self.assertEqual(ws["B2"].value.strftime("%d.%m.%Y"), "30.09.2026")
            self.assertEqual(ws["C2"].value, "СН-041285")
            self.assertEqual(ws["D2"].value.strftime("%d.%m.%Y %H:%M"), "29.09.2026 10:45")
            self.assertEqual(ws["E2"].value.strftime("%d.%m.%Y %H:%M"), "29.09.2026 10:49")
            self.assertEqual(ws["F2"].value.strftime("%d.%m.%Y %H:%M"), "29.09.2026 10:49")
            self.assertIsNone(ws["G2"].value)
            self.assertEqual(ws["H2"].value.strftime("%d.%m.%Y %H:%M"), "29.09.2026 10:50")
            self.assertEqual(ws["I2"].value.strftime("%d.%m.%Y %H:%M"), "29.09.2026 10:51")
            self.assertEqual(ws["J2"].value.strftime("%d.%m.%Y %H:%M"), "29.09.2026 11:00")
            self.assertEqual(ws["K2"].value.strftime("%d.%m.%Y %H:%M"), "29.09.2026 11:10")
            self.assertEqual(ws["L2"].value, "Нет допуска к месту работ")
            self.assertEqual(ws["M2"].value.strftime("%d.%m.%Y %H:%M"), "29.09.2026 11:20")
            self.assertEqual(ws["N2"].value.strftime("%d.%m.%Y %H:%M"), "29.09.2026 11:25")
            self.assertIsNone(ws["O2"].value)
            self.assertEqual(ws["P2"].value.strftime("%d.%m.%Y %H:%M"), "29.09.2026 11:30")
            self.assertEqual(ws["Q2"].value.strftime("%d.%m.%Y %H:%M"), "29.09.2026 11:35")
            self.assertEqual(ws["R2"].value.strftime("%d.%m.%Y %H:%M"), "29.09.2026 11:40")
            self.assertIsNone(ws["S2"].value)
            self.assertIsNone(ws["T2"].value)
            self.assertIsNone(ws["U2"].value)

            self.assertEqual(ws["B2"].number_format, "dd.mm.yyyy")
            self.assertEqual(ws["D2"].number_format, "dd.mm.yyyy hh:mm")
            self.assertEqual(ws["C2"].number_format, "@")
            self.assertEqual(wb["Инструкция"]["A1"].value, "Шаблон загрузки данных из MAX в Платформу РПО")

    def test_max_template_limit_is_explicit(self) -> None:
        record = self._record()
        with tempfile.TemporaryDirectory() as tmp:
            with self.assertRaisesRegex(ValueError, "максимум"):
                build_export([record] * 501, tmp, 30)


if __name__ == "__main__":
    unittest.main()
