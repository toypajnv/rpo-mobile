from __future__ import annotations

import json
from collections import Counter
from datetime import datetime, timedelta, timezone
from io import BytesIO
from pathlib import Path
from typing import Any

from fastapi import APIRouter, Depends, Form, HTTPException, Request
from fastapi.responses import RedirectResponse, StreamingResponse
from fastapi.templating import Jinja2Templates
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from pydantic import BaseModel, Field
from sqlalchemy import DateTime, String, Text, func, select
from sqlalchemy.orm import Mapped, Session, mapped_column

from .database import Base, get_db
from .security import verify_password

BASE_DIR = Path(__file__).resolve().parent
templates = Jinja2Templates(directory=BASE_DIR / "templates")
router = APIRouter(prefix="/opros", tags=["Опрос ПБ"])

SURVEY_VERSION = "2026-10"
SURVEY_ADMIN_USERNAME = "admin"
SURVEY_ADMIN_PASSWORD_HASH = "$argon2id$v=19$m=65536,t=3,p=4$N5zlhe6Ioo7FXDWeyh5Ydg$jZPjn68fguEA0x80AhLqadCREH+EHToFsTCNxSvNUPg"

QUESTION_LABELS = {
    "q1": "Если Вы видите, что работник подрядной организации выполняет работу с нарушением требований ПБ, что Вы обычно делаете?",
    "q2": "Как часто Вы сталкиваетесь с нарушениями требований ПБ со стороны подрядчиков?",
    "q3": "Насколько Вы считаете, что должны лично вмешаться, если видите нарушение ПБ у подрядчика?",
    "q4": "Что чаще всего мешает Вам вмешаться в работу подрядчика?",
    "q5": "Представьте: Вы проходите мимо места работы подрядчика и видите явное нарушение ПБ. Ваш руководитель находится далеко. Что Вы будете делать?",
    "q6": "Знаете ли Вы установленный порядок действий, если при работе подрядчика выявлено нарушение ПБ?",
    "q7": "Как обычно реагирует Ваш непосредственный руководитель, если Вы сообщаете о нарушении ПБ у подрядчика?",
    "q8": "Как Вы думаете, что произойдёт, если Вы остановите работу подрядчика из-за нарушения ПБ?",
    "q9": "Что помогло бы Вам чаще вмешиваться при нарушениях ПБ у подрядчиков?",
    "q10": "Вспомните последний случай, когда Вы видели нарушение ПБ у подрядчика. Что Вы сделали?",
    "q11": "Что, на Ваш взгляд, чаще всего заставляет работников проходить мимо нарушения?",
    "q12": "Что нужно изменить, чтобы работники действительно не проходили мимо нарушений ПБ подрядчиков?",
    "q13": "Если я вижу нарушение требований ПБ при работе подрядчика, я считаю своей обязанностью отреагировать, даже если это не мой участок и не мой работник.",
    "q14": "Ваши предложения по процессу остановки работ",
}

SINGLE_OPTIONS = {
    "q1": [
        "Сразу останавливаю работу / прошу прекратить опасное действие",
        "Делаю замечание, но работу не останавливаю",
        "Сообщаю своему руководителю",
        "Сообщаю руководителю подрядчика",
        "Прохожу мимо, если нарушение не кажется серьёзным",
        "Ничего не делаю",
        "Другое",
    ],
    "q2": [
        "Практически каждый день",
        "Несколько раз в неделю",
        "Несколько раз в месяц",
        "Редко",
        "Практически никогда",
    ],
    "q3": ["1", "2", "3", "4", "5"],
    "q5": [
        "Сам остановлю работу",
        "Сделаю замечание и попрошу устранить нарушение",
        "Позвоню своему руководителю",
        "Сообщу руководителю подрядчика",
        "Подожду, пока придёт ответственный",
        "Скорее всего, не буду вмешиваться",
        "Другое",
    ],
    "q6": ["Да, хорошо знаю", "Примерно знаю", "Не знаю"],
    "q7": [
        "Поддерживает и помогает устранить нарушение",
        "Спокойно принимает информацию",
        "Просит сначала разобраться самостоятельно",
        "Считает, что этим должен заниматься подрядчик",
        "Негативно реагирует",
        "Случаев не было",
        "Затрудняюсь ответить",
    ],
    "q8": [
        "Руководитель поддержит",
        "Нарушение устранят и работу продолжат",
        "Возникнут споры с подрядчиком",
        "Мне скажут, что я вмешиваюсь не в свою работу",
        "Работа будет необоснованно задержана",
        "Не знаю",
        "Другое",
    ],
    "q13": [
        "Полностью согласен",
        "Скорее согласен",
        "Скорее не согласен",
        "Не согласен",
        "Затрудняюсь ответить",
    ],
}

MULTI_OPTIONS = {
    "q4": [
        "Считаю, что это ответственность руководителя подрядчика",
        "Считаю, что это не моя зона ответственности",
        "Не хочу конфликтовать с подрядчиком",
        "Не уверен, что нарушение действительно является нарушением",
        "Не знаю, к кому обращаться",
        "Боюсь негативной реакции руководителя",
        "Боюсь негативной реакции подрядчика",
        "Не хочу задерживать выполнение работ",
        "Думаю, что кто-то другой уже обратит внимание",
        "Нарушение кажется незначительным",
        "Ничего не мешает — я вмешиваюсь",
        "Другое",
    ],
    "q9": [
        "Чётко закреплённое право/обязанность вмешиваться",
        "Поддержка руководителя",
        "Понятный порядок действий",
        "Обучение на реальных ситуациях",
        "Возможность быстро сообщить о нарушении без бюрократии",
        "Положительная оценка таких действий",
        "Обратная связь о результатах",
        "Личный пример руководителей",
        "Ничего — я и сейчас вмешиваюсь",
        "Другое",
    ],
}

REQUIRED_KEYS = {"q1", "q2", "q3", "q4", "q5", "q6", "q7", "q8", "q9", "q13"}
OPEN_KEYS = {"q10", "q11", "q12", "q14"}


def utcnow() -> datetime:
    return datetime.now(timezone.utc)


def _as_utc(value: datetime | None) -> datetime | None:
    if value is None:
        return None
    if value.tzinfo is None:
        return value.replace(tzinfo=timezone.utc)
    return value.astimezone(timezone.utc)


class SafetySurveyResponse(Base):
    __tablename__ = "safety_survey_responses"

    id: Mapped[int] = mapped_column(primary_key=True, autoincrement=True)
    survey_version: Mapped[str] = mapped_column(String(32), default=SURVEY_VERSION, index=True)
    answers_json: Mapped[str] = mapped_column(Text, nullable=False)
    submitted_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=utcnow, index=True)


class SurveyPayload(BaseModel):
    answers: dict[str, Any] = Field(default_factory=dict)


def _normalize_other(value: str) -> str:
    value = " ".join((value or "").strip().split())
    if value.startswith("Другое:"):
        return "Другое"
    return value


def _clean_text(value: Any, limit: int = 4000) -> str:
    text = str(value or "").strip()
    return text[:limit]


def _validate_answers(raw: dict[str, Any]) -> dict[str, Any]:
    if not isinstance(raw, dict):
        raise HTTPException(status_code=422, detail="Некорректный формат ответов")

    cleaned: dict[str, Any] = {}

    for key, options in SINGLE_OPTIONS.items():
        value = _clean_text(raw.get(key), 1200)
        normalized = _normalize_other(value)
        if key in REQUIRED_KEYS and not value:
            raise HTTPException(status_code=422, detail=f"Не заполнен обязательный вопрос {key}")
        if value and normalized not in options:
            raise HTTPException(status_code=422, detail=f"Некорректный ответ {key}")
        cleaned[key] = value

    for key, options in MULTI_OPTIONS.items():
        value = raw.get(key)
        if not isinstance(value, list):
            raise HTTPException(status_code=422, detail=f"Некорректный ответ {key}")
        items: list[str] = []
        for item in value:
            text = _clean_text(item, 1200)
            if not text:
                continue
            if _normalize_other(text) not in options:
                raise HTTPException(status_code=422, detail=f"Некорректный ответ {key}")
            if text not in items:
                items.append(text)
        if key in REQUIRED_KEYS and not items:
            raise HTTPException(status_code=422, detail=f"Не заполнен обязательный вопрос {key}")
        if key == "q9" and len(items) > 3:
            raise HTTPException(status_code=422, detail="В вопросе 9 можно выбрать не более 3 вариантов")
        cleaned[key] = items

    for key in OPEN_KEYS:
        cleaned[key] = _clean_text(raw.get(key), 4000)

    return cleaned


def _is_survey_admin(request: Request) -> bool:
    return request.session.get("survey_admin") is True


def _require_survey_admin(request: Request) -> None:
    if not _is_survey_admin(request):
        raise HTTPException(status_code=401, detail="Требуется вход")


def _answer_dict(row: SafetySurveyResponse) -> dict[str, Any]:
    try:
        value = json.loads(row.answers_json)
        return value if isinstance(value, dict) else {}
    except Exception:
        return {}


def _distribution(rows: list[SafetySurveyResponse], key: str, *, multi: bool = False) -> list[dict[str, Any]]:
    counter: Counter[str] = Counter()
    for row in rows:
        answers = _answer_dict(row)
        value = answers.get(key, [] if multi else "")
        values = value if multi and isinstance(value, list) else [value]
        for item in values:
            label = _normalize_other(_clean_text(item, 1200))
            if label:
                counter[label] += 1
    total = len(rows) or 1
    return [
        {
            "label": label,
            "count": count,
            "percent": round(count * 100 / total, 1),
        }
        for label, count in counter.most_common()
    ]


def _readiness_index(rows: list[SafetySurveyResponse]) -> int:
    if not rows:
        return 0

    scores: list[float] = []
    q1_map = {
        "Сразу останавливаю работу / прошу прекратить опасное действие": 100,
        "Делаю замечание, но работу не останавливаю": 75,
        "Сообщаю своему руководителю": 50,
        "Сообщаю руководителю подрядчика": 50,
        "Прохожу мимо, если нарушение не кажется серьёзным": 25,
        "Ничего не делаю": 0,
    }
    q5_map = {
        "Сам остановлю работу": 100,
        "Сделаю замечание и попрошу устранить нарушение": 80,
        "Позвоню своему руководителю": 60,
        "Сообщу руководителю подрядчика": 60,
        "Подожду, пока придёт ответственный": 30,
        "Скорее всего, не буду вмешиваться": 0,
    }
    q6_map = {"Да, хорошо знаю": 100, "Примерно знаю": 60, "Не знаю": 0}
    q13_map = {
        "Полностью согласен": 100,
        "Скорее согласен": 75,
        "Скорее не согласен": 25,
        "Не согласен": 0,
    }

    for row in rows:
        answers = _answer_dict(row)
        local: list[float] = []

        q1 = _normalize_other(_clean_text(answers.get("q1")))
        if q1 in q1_map:
            local.append(q1_map[q1])

        q3 = _clean_text(answers.get("q3"))
        if q3 in {"1", "2", "3", "4", "5"}:
            local.append((int(q3) - 1) * 25)

        q5 = _normalize_other(_clean_text(answers.get("q5")))
        if q5 in q5_map:
            local.append(q5_map[q5])

        q6 = _clean_text(answers.get("q6"))
        if q6 in q6_map:
            local.append(q6_map[q6])

        q13 = _clean_text(answers.get("q13"))
        if q13 in q13_map:
            local.append(q13_map[q13])

        if local:
            scores.append(sum(local) / len(local))

    return round(sum(scores) / len(scores)) if scores else 0


def _stats_payload(rows: list[SafetySurveyResponse]) -> dict[str, Any]:
    total = len(rows)
    cutoff = utcnow() - timedelta(hours=24)
    last_24h = sum(1 for row in rows if _as_utc(row.submitted_at) and _as_utc(row.submitted_at) >= cutoff)

    q5_stop = 0
    q6_know = 0
    q13_agree = 0
    for row in rows:
        answers = _answer_dict(row)
        if answers.get("q5") == "Сам остановлю работу":
            q5_stop += 1
        if answers.get("q6") == "Да, хорошо знаю":
            q6_know += 1
        if answers.get("q13") in {"Полностью согласен", "Скорее согласен"}:
            q13_agree += 1

    def pct(value: int) -> int:
        return round(value * 100 / total) if total else 0

    daily_counter: Counter[str] = Counter()
    for row in rows:
        if row.submitted_at:
            daily_counter[row.submitted_at.date().isoformat()] += 1

    distributions: dict[str, list[dict[str, Any]]] = {}
    for key in SINGLE_OPTIONS:
        distributions[key] = _distribution(rows, key)
    for key in MULTI_OPTIONS:
        distributions[key] = _distribution(rows, key, multi=True)

    open_answers: dict[str, list[dict[str, Any]]] = {key: [] for key in OPEN_KEYS}
    for row in rows:
        answers = _answer_dict(row)
        for key in OPEN_KEYS:
            text = _clean_text(answers.get(key), 4000)
            if text:
                open_answers[key].append(
                    {
                        "id": row.id,
                        "submitted_at": row.submitted_at.isoformat() if row.submitted_at else "",
                        "answer": text,
                    }
                )

    return {
        "total": total,
        "last_24h": last_24h,
        "readiness_index": _readiness_index(rows),
        "direct_stop_percent": pct(q5_stop),
        "know_order_percent": pct(q6_know),
        "duty_agree_percent": pct(q13_agree),
        "top_barriers": distributions["q4"][:6],
        "top_enablers": distributions["q9"][:6],
        "distributions": distributions,
        "question_labels": QUESTION_LABELS,
        "open_answers": open_answers,
        "daily": [{"date": date, "count": count} for date, count in sorted(daily_counter.items())],
    }


@router.get("/")
def survey_page(request: Request):
    return templates.TemplateResponse(
        "survey.html",
        {"request": request, "question_labels": QUESTION_LABELS},
        headers={"Cache-Control": "no-store"},
    )


@router.post("/api/responses")
def submit_survey(payload: SurveyPayload, db: Session = Depends(get_db)):
    answers = _validate_answers(payload.answers)
    row = SafetySurveyResponse(
        survey_version=SURVEY_VERSION,
        answers_json=json.dumps(answers, ensure_ascii=False),
    )
    db.add(row)
    db.commit()
    db.refresh(row)
    return {"ok": True, "id": row.id}


@router.get("/statistika")
def survey_statistics_page(request: Request):
    return templates.TemplateResponse(
        "survey_stats.html",
        {"request": request, "logged_in": _is_survey_admin(request)},
        headers={"Cache-Control": "no-store"},
    )


@router.post("/statistika/login")
def survey_statistics_login(
    request: Request,
    username: str = Form(...),
    password: str = Form(...),
):
    if username == SURVEY_ADMIN_USERNAME and verify_password(password, SURVEY_ADMIN_PASSWORD_HASH):
        request.session["survey_admin"] = True
        return RedirectResponse("/opros/statistika", status_code=303)
    return templates.TemplateResponse(
        "survey_stats.html",
        {
            "request": request,
            "logged_in": False,
            "login_error": "Неверный логин или пароль",
        },
        status_code=401,
        headers={"Cache-Control": "no-store"},
    )


@router.post("/statistika/logout")
def survey_statistics_logout(request: Request):
    request.session.pop("survey_admin", None)
    return RedirectResponse("/opros/statistika", status_code=303)


@router.get("/api/stats")
def survey_statistics(request: Request, db: Session = Depends(get_db)):
    _require_survey_admin(request)
    rows = list(
        db.scalars(
            select(SafetySurveyResponse)
            .where(SafetySurveyResponse.survey_version == SURVEY_VERSION)
            .order_by(SafetySurveyResponse.submitted_at.desc(), SafetySurveyResponse.id.desc())
        )
    )
    return _stats_payload(rows)


@router.get("/api/export.xlsx")
def survey_export(request: Request, db: Session = Depends(get_db)):
    _require_survey_admin(request)
    rows = list(
        db.scalars(
            select(SafetySurveyResponse)
            .where(SafetySurveyResponse.survey_version == SURVEY_VERSION)
            .order_by(SafetySurveyResponse.submitted_at.asc(), SafetySurveyResponse.id.asc())
        )
    )
    stats = _stats_payload(rows)

    wb = Workbook()
    ws = wb.active
    ws.title = "Ответы"
    headers = ["ID", "Дата/время (UTC)"] + [f"{key.upper()}. {QUESTION_LABELS[key]}" for key in QUESTION_LABELS]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9EAF7")
        cell.alignment = Alignment(wrap_text=True, vertical="top")

    for row in rows:
        answers = _answer_dict(row)
        values = [row.id, row.submitted_at.isoformat() if row.submitted_at else ""]
        for key in QUESTION_LABELS:
            value = answers.get(key, "")
            if isinstance(value, list):
                value = "; ".join(str(v) for v in value)
            values.append(value)
        ws.append(values)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.column_dimensions["A"].width = 8
    ws.column_dimensions["B"].width = 24
    for column_letter in "CDEFGHIJKLMNOP":
        ws.column_dimensions[column_letter].width = 32

    summary = wb.create_sheet("Сводка")
    summary.append(["Показатель", "Значение"])
    summary.append(["Количество участников", stats["total"]])
    summary.append(["Ответов за 24 часа", stats["last_24h"]])
    summary.append(["Индекс готовности вмешаться", stats["readiness_index"]])
    summary.append(["Готовы самостоятельно остановить работу, %", stats["direct_stop_percent"]])
    summary.append(["Хорошо знают порядок действий, %", stats["know_order_percent"]])
    summary.append(["Считают обязанностью отреагировать, %", stats["duty_agree_percent"]])
    summary.append([])
    summary.append(["Основные барьеры", "Количество"])
    for item in stats["top_barriers"]:
        summary.append([item["label"], item["count"]])
    summary.append([])
    summary.append(["Что поможет чаще вмешиваться", "Количество"])
    for item in stats["top_enablers"]:
        summary.append([item["label"], item["count"]])

    for cell in summary[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9EAF7")
    summary.column_dimensions["A"].width = 68
    summary.column_dimensions["B"].width = 18

    stream = BytesIO()
    wb.save(stream)
    stream.seek(0)
    filename = f"safety-survey-{utcnow().date().isoformat()}.xlsx"
    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )
