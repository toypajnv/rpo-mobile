from __future__ import annotations

import hashlib
import json
import mimetypes
import re
import secrets
import shutil
from datetime import datetime, timedelta, timezone
from pathlib import Path
from typing import Any

from fastapi import APIRouter, Depends, File, Form, HTTPException, Request, UploadFile
from fastapi.responses import FileResponse, HTMLResponse, JSONResponse, RedirectResponse
from fastapi.templating import Jinja2Templates
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from pydantic import BaseModel, Field
from sqlalchemy import and_, func, or_, select
from sqlalchemy.orm import Session

from .config import settings
from .database import get_db
from .models import Operator
from .pb_mng_models import (
    PbAccessRestriction,
    PbAction,
    PbEmailRoute,
    PbPhoto,
    PbStop,
    PbViolationRule,
)
from .services.mailer import send_notification
from .stop_registry import StopRegistryRecord, normalize_pass_number

BASE_DIR = Path(__file__).resolve().parent
PB_DIR = BASE_DIR / "pb_mng"
PB_UPLOAD_DIR = Path(getattr(settings, "pb_mng_upload_dir", "./data/pb_mng/photos"))
PB_EXPORT_DIR = Path(getattr(settings, "pb_mng_export_dir", "./data/pb_mng/exports"))
PB_UPLOAD_DIR.mkdir(parents=True, exist_ok=True)
PB_EXPORT_DIR.mkdir(parents=True, exist_ok=True)

templates = Jinja2Templates(directory=BASE_DIR / "templates")
router = APIRouter()

SEVERITY_LABELS = {"gross": "Грубое", "significant": "Значительное", "minor": "Незначительное"}
STATUS_LABELS = {
    "pending_verification": "На проверке координатора",
    "returned_for_revision": "Возвращена на доработку",
    "rejected": "Отклонена",
    "awaiting_resolution": "Ожидается устранение",
    "resolution_submitted": "Устранение заявлено",
    "resolution_revision": "Устранение требует доработки",
    "awaiting_pkm": "Ожидается ПКМ",
    "awaiting_training": "Ожидается обучение",
    "ready_for_unblock": "Готово к снятию блокировки",
    "closed": "Закрыта",
}


def utcnow() -> datetime:
    return datetime.now(timezone.utc)


def _as_utc(value: datetime | None) -> datetime | None:
    if value is None:
        return None
    if value.tzinfo is None:
        return value.replace(tzinfo=timezone.utc)
    return value.astimezone(timezone.utc)


def _json(value: str | None, default: Any) -> Any:
    try:
        parsed = json.loads(value or "")
        return parsed
    except Exception:
        return default


def _token_hash(value: str) -> str:
    return hashlib.sha256((value or "").encode("utf-8")).hexdigest()


def _clean(value: Any, limit: int = 1000) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip()[:limit]


def _device_headers(request: Request) -> tuple[str, str]:
    device_id = _clean(request.headers.get("X-PB-Device"), 160)
    token = _clean(request.headers.get("X-PB-Token"), 300)
    if not device_id or not token:
        raise HTTPException(status_code=401, detail="Устройство не идентифицировано")
    return device_id, _token_hash(token)


def _current_operator_or_none(request: Request, db: Session) -> Operator | None:
    operator_id = request.session.get("operator_id")
    if not operator_id:
        return None
    operator = db.get(Operator, operator_id)
    if not operator or not operator.is_active:
        request.session.clear()
        return None
    return operator


def coordinator_operator(request: Request, db: Session = Depends(get_db)) -> Operator:
    operator = _current_operator_or_none(request, db)
    if not operator:
        raise HTTPException(status_code=401, detail="Требуется вход координатора")
    if request.method.upper() != "GET" and (getattr(operator, "role", "operator") or "operator") == "manager":
        raise HTTPException(status_code=403, detail="Роль «Руководитель» работает только в режиме просмотра")
    return operator


def _seed_catalog(db: Session) -> None:
    if db.scalar(select(func.count(PbViolationRule.id))) or 0:
        return
    seed_file = PB_DIR / "classifier.json.b64"
    if not seed_file.exists():
        return
    import base64
    import zlib

    payload = json.loads(zlib.decompress(base64.b64decode(seed_file.read_text(encoding="utf-8"))).decode("utf-8"))
    rows = []
    for item in payload.get("items", []):
        rows.append(
            PbViolationRule(
                id=str(item.get("id", ""))[:80],
                source=str(item.get("source", ""))[:40],
                number=int(item.get("number") or 0),
                barrier=str(item.get("barrier", "")),
                text=str(item.get("text", "")),
                severities_json=json.dumps(item.get("severities") or [], ensure_ascii=False),
                criteria_json=json.dumps(item.get("criteria") or {}, ensure_ascii=False),
                requires_context=bool(item.get("requires_context")),
                active=True,
            )
        )
    db.add_all(rows)
    db.commit()


def _reference_seed() -> dict:
    path = PB_DIR / "catalog_seed.json"
    if not path.exists():
        return {}
    return json.loads(path.read_text(encoding="utf-8"))


def _rule_out(rule: PbViolationRule) -> dict:
    return {
        "id": rule.id,
        "source": rule.source,
        "number": rule.number,
        "barrier": rule.barrier,
        "text": rule.text,
        "severities": _json(rule.severities_json, []),
        "criteria": _json(rule.criteria_json, {}),
        "requires_context": bool(rule.requires_context),
    }


def _effective_severity(stop: PbStop, db: Session, *, record_action: bool = True) -> str:
    if stop.current_severity != "minor":
        return stop.current_severity
    if stop.status in {"rejected", "closed", "resolution_submitted", "ready_for_unblock"}:
        return stop.current_severity
    occurred_at = _as_utc(stop.occurred_at) or utcnow()
    if utcnow() < occurred_at + timedelta(minutes=60):
        return stop.current_severity
    stop.current_severity = "significant"
    stop.updated_at = utcnow()
    if record_action:
        _audit(db, stop, "system", "Система", "minor_escalated", {"rule": "minor_over_60_minutes"})
    db.commit()
    return stop.current_severity


def _audit(db: Session, stop: PbStop, actor_type: str, actor_name: str, action: str, data: dict | None = None) -> None:
    db.add(
        PbAction(
            stop_id=stop.id,
            actor_type=actor_type,
            actor_name=_clean(actor_name, 240),
            action=action[:80],
            data_json=json.dumps(data or {}, ensure_ascii=False),
        )
    )


def _photo_out(photo: PbPhoto) -> dict:
    return {
        "id": photo.id,
        "phase": photo.phase,
        "url": f"/api/pb-mng/photos/{photo.id}",
        "original_name": photo.original_name or "",
        "mime_type": photo.mime_type or "",
        "created_at": photo.created_at.isoformat(),
    }


def _stop_out(stop: PbStop, db: Session, *, include_actions: bool = False) -> dict:
    _effective_severity(stop, db)
    active_restrictions = [r for r in stop.restrictions if r.active and (r.ends_at is None or (_as_utc(r.ends_at) or utcnow()) > utcnow())]
    result = {
        "id": stop.public_id,
        "initiator": {"name": stop.initiator_name, "unit": stop.initiator_unit, "pass": stop.initiator_pass},
        "occurred_at": stop.occurred_at.isoformat(),
        "created_at": stop.created_at.isoformat(),
        "updated_at": stop.updated_at.isoformat(),
        "block": stop.block,
        "structural_unit": stop.structural_unit,
        "field": stop.field_name,
        "location": stop.location,
        "contractor": stop.contractor,
        "subcontractor": stop.subcontractor,
        "work_type": stop.work_type,
        "permit_number": stop.permit_number,
        "violation": {
            "id": stop.violation_id,
            "text": stop.violation_text,
            "barrier": stop.violation_barrier,
        },
        "initial_severity": stop.initial_severity,
        "current_severity": stop.current_severity,
        "severity_label": SEVERITY_LABELS.get(stop.current_severity, stop.current_severity),
        "severity_context": stop.severity_context,
        "description": stop.description,
        "responsible": {"fio": stop.responsible_fio, "position": stop.responsible_position, "pass": stop.responsible_pass},
        "crew": _json(stop.crew_json, []),
        "supervisor": {"fio": stop.supervisor_fio, "pass": stop.supervisor_pass},
        "status": stop.status,
        "status_label": STATUS_LABELS.get(stop.status, stop.status),
        "verification_note": stop.verification_note,
        "measures": _json(stop.measures_json, []),
        "course_name": stop.course_name,
        "course_status": stop.course_status,
        "pkm_required": stop.pkm_required,
        "pkm_status": stop.pkm_status,
        "pkm_text": stop.pkm_text,
        "resolution_comment": stop.resolution_comment,
        "resolution_submitted_at": stop.resolution_submitted_at.isoformat() if stop.resolution_submitted_at else None,
        "verified_at": stop.verified_at.isoformat() if stop.verified_at else None,
        "closed_at": stop.closed_at.isoformat() if stop.closed_at else None,
        "photos": [_photo_out(x) for x in sorted(stop.photos, key=lambda p: p.created_at)],
        "restrictions": [
            {
                "id": r.id,
                "pass_number": r.pass_number,
                "fio": r.fio,
                "kind": r.restriction_kind,
                "ends_at": r.ends_at.isoformat() if r.ends_at else None,
            }
            for r in active_restrictions
        ],
        "mail_status": stop.mail_status,
    }
    if include_actions:
        result["actions"] = [
            {
                "at": a.created_at.isoformat(),
                "actor_type": a.actor_type,
                "actor_name": a.actor_name,
                "action": a.action,
                "data": _json(a.data_json, {}),
            }
            for a in sorted(stop.actions, key=lambda x: x.created_at)
        ]
    return result


def _find_own_stop(db: Session, public_id: str, request: Request) -> PbStop:
    device_id, token_hash = _device_headers(request)
    stop = db.scalar(
        select(PbStop).where(
            PbStop.public_id == public_id,
            PbStop.device_id == device_id,
            PbStop.device_token_hash == token_hash,
        )
    )
    if not stop:
        raise HTTPException(status_code=404, detail="Остановка не найдена")
    return stop


def _save_photos(db: Session, stop: PbStop, uploads: list[UploadFile], phase: str) -> int:
    saved = 0
    allowed = {"image/jpeg": ".jpg", "image/png": ".png", "image/webp": ".webp", "image/heic": ".heic", "image/heif": ".heif"}
    target_dir = PB_UPLOAD_DIR / str(stop.id)
    target_dir.mkdir(parents=True, exist_ok=True)
    for upload in uploads[:5]:
        mime = (upload.content_type or "").lower()
        if mime not in allowed:
            continue
        raw = upload.file.read(10 * 1024 * 1024 + 1)
        if len(raw) > 10 * 1024 * 1024:
            raise HTTPException(status_code=413, detail="Фото превышает 10 МБ")
        stored_name = f"{phase}_{secrets.token_hex(16)}{allowed[mime]}"
        path = target_dir / stored_name
        path.write_bytes(raw)
        db.add(PbPhoto(stop_id=stop.id, phase=phase, stored_name=f"{stop.id}/{stored_name}", original_name=_clean(upload.filename, 255), mime_type=mime))
        saved += 1
    return saved


def _new_public_id(db: Session) -> str:
    year = utcnow().year
    prefix = f"ОР-{year}-"
    latest = db.scalar(select(PbStop.public_id).where(PbStop.public_id.like(f"{prefix}%")).order_by(PbStop.id.desc()).limit(1))
    seq = 1
    if latest:
        try:
            seq = int(latest.rsplit("-", 1)[1]) + 1
        except Exception:
            seq = (db.scalar(select(func.count(PbStop.id)).where(PbStop.public_id.like(f"{prefix}%"))) or 0) + 1
    return f"{prefix}{seq:05d}"


def _resolve_severity(rule: PbViolationRule, requested: str, context: str) -> str:
    severities = _json(rule.severities_json, [])
    if not severities:
        raise HTTPException(status_code=400, detail="Для нарушения не настроена категория")
    if len(severities) == 1:
        return severities[0]
    if requested not in severities:
        raise HTTPException(status_code=400, detail="Выберите применимую категорию по критерию классификатора")
    if rule.requires_context and not _clean(context, 2000):
        raise HTTPException(status_code=400, detail="Укажите основание выбора категории")
    return requested


class ReviewPayload(BaseModel):
    action: str
    note: str = ""
    severity: str = ""
    measures: list[str] = Field(default_factory=list)
    course_name: str = ""
    pkm_required: bool = False
    block_responsible: bool = False
    block_days: int | None = Field(default=None, ge=1, le=365)


def _next_after_conditions(stop: PbStop) -> str:
    if stop.pkm_required and stop.pkm_status != "accepted":
        return "awaiting_pkm"
    if stop.course_status == "required":
        return "awaiting_training"
    has_active = any(r.active and (r.ends_at is None or (_as_utc(r.ends_at) or utcnow()) > utcnow()) for r in stop.restrictions)
    return "ready_for_unblock" if has_active else "closed"


class ResolutionReviewPayload(BaseModel):
    action: str
    note: str = ""


class PkmPayload(BaseModel):
    action: str
    text: str = ""


class EmailRoutePayload(BaseModel):
    block: str = ""
    contractor: str = ""
    role: str
    recipient_name: str = ""
    email: str


@router.get("/api/pb-mng/catalog")
def pb_catalog(db: Session = Depends(get_db)):
    _seed_catalog(db)
    rules = db.scalars(select(PbViolationRule).where(PbViolationRule.active.is_(True)).order_by(PbViolationRule.source, PbViolationRule.number)).all()
    contractors = [x for x in db.scalars(select(StopRegistryRecord.company).where(StopRegistryRecord.company != "").distinct().order_by(StopRegistryRecord.company)).all() if x]
    ref = _reference_seed()
    return {
        **ref,
        "contractors": contractors,
        "violations": [_rule_out(rule) for rule in rules],
        "status_labels": STATUS_LABELS,
    }


@router.post("/api/pb-mng/stops", status_code=201)
def create_stop(
    request: Request,
    payload_json: str = Form(...),
    photos: list[UploadFile] = File(default=[]),
    db: Session = Depends(get_db),
):
    _seed_catalog(db)
    try:
        payload = json.loads(payload_json)
    except Exception as exc:
        raise HTTPException(status_code=400, detail="Некорректная карточка остановки") from exc

    device_id = _clean(payload.get("device_id"), 160)
    device_token = _clean(payload.get("device_token"), 300)
    if not device_id or not device_token:
        raise HTTPException(status_code=400, detail="Не задан идентификатор устройства")
    initiator = payload.get("initiator") or {}
    initiator_name = _clean(initiator.get("name"), 240)
    if not initiator_name:
        raise HTTPException(status_code=400, detail="Укажите ФИО инициатора")

    violation_id = _clean(payload.get("violation_id"), 80)
    rule = db.get(PbViolationRule, violation_id)
    if not rule or not rule.active:
        raise HTTPException(status_code=400, detail="Выберите нарушение из классификатора")
    severity_context = _clean(payload.get("severity_context"), 2000)
    severity = _resolve_severity(rule, _clean(payload.get("severity"), 30), severity_context)

    try:
        occurred_at = datetime.fromisoformat(str(payload.get("occurred_at") or "").replace("Z", "+00:00"))
        if occurred_at.tzinfo is None:
            occurred_at = occurred_at.replace(tzinfo=timezone.utc)
    except Exception:
        occurred_at = utcnow()

    responsible = payload.get("responsible") or {}
    stop = PbStop(
        public_id=_new_public_id(db),
        device_id=device_id,
        device_token_hash=_token_hash(device_token),
        initiator_name=initiator_name,
        initiator_unit=_clean(initiator.get("unit"), 160),
        initiator_pass=_clean(initiator.get("pass"), 60),
        occurred_at=occurred_at,
        block=_clean(payload.get("block"), 160),
        structural_unit=_clean(payload.get("structural_unit"), 160),
        field_name=_clean(payload.get("field"), 180),
        location=_clean(payload.get("location"), 300),
        contractor=_clean(payload.get("contractor"), 300),
        subcontractor=_clean(payload.get("subcontractor"), 300),
        work_type=_clean(payload.get("work_type"), 180),
        permit_number=_clean(payload.get("permit_number"), 100),
        violation_id=rule.id,
        violation_text=rule.text,
        violation_barrier=rule.barrier,
        initial_severity=severity,
        current_severity=severity,
        severity_context=severity_context,
        description=_clean(payload.get("description"), 5000),
        responsible_fio=_clean(responsible.get("fio"), 240),
        responsible_position=_clean(responsible.get("position"), 240),
        responsible_pass=normalize_pass_number(responsible.get("pass")) or _clean(responsible.get("pass"), 60),
        crew_json=json.dumps(payload.get("crew") or [], ensure_ascii=False),
        supervisor_fio=_clean((payload.get("supervisor") or {}).get("fio"), 240),
        supervisor_pass=normalize_pass_number((payload.get("supervisor") or {}).get("pass")) or _clean((payload.get("supervisor") or {}).get("pass"), 60),
        status="pending_verification",
    )
    db.add(stop)
    db.flush()
    count = _save_photos(db, stop, photos, "violation")
    if count == 0:
        raise HTTPException(status_code=400, detail="Добавьте хотя бы одно фото нарушения")
    _audit(db, stop, "worker", stop.initiator_name, "created", {"photos": count})
    db.commit()
    db.refresh(stop)
    return _stop_out(stop, db, include_actions=True)


@router.get("/api/pb-mng/my-stops")
def my_stops(request: Request, db: Session = Depends(get_db)):
    device_id, token_hash = _device_headers(request)
    rows = db.scalars(select(PbStop).where(PbStop.device_id == device_id, PbStop.device_token_hash == token_hash).order_by(PbStop.created_at.desc()).limit(200)).all()
    return {"items": [_stop_out(x, db) for x in rows]}


@router.get("/api/pb-mng/stops/{public_id}")
def my_stop(public_id: str, request: Request, db: Session = Depends(get_db)):
    stop = _find_own_stop(db, public_id, request)
    return _stop_out(stop, db, include_actions=True)


@router.post("/api/pb-mng/stops/{public_id}/resubmit")
def resubmit_stop(public_id: str, request: Request, payload_json: str = Form(...), photos: list[UploadFile] = File(default=[]), db: Session = Depends(get_db)):
    stop = _find_own_stop(db, public_id, request)
    if stop.status != "returned_for_revision":
        raise HTTPException(status_code=409, detail="Карточка сейчас не находится на доработке")
    payload = json.loads(payload_json or "{}")
    stop.description = _clean(payload.get("description", stop.description), 5000)
    stop.location = _clean(payload.get("location", stop.location), 300)
    stop.status = "pending_verification"
    stop.verification_note = ""
    stop.updated_at = utcnow()
    count = _save_photos(db, stop, photos, "violation") if photos else 0
    _audit(db, stop, "worker", stop.initiator_name, "resubmitted", {"new_photos": count})
    db.commit()
    return _stop_out(stop, db, include_actions=True)


@router.post("/api/pb-mng/stops/{public_id}/resolution")
def submit_resolution(public_id: str, request: Request, comment: str = Form(""), photos: list[UploadFile] = File(default=[]), db: Session = Depends(get_db)):
    stop = _find_own_stop(db, public_id, request)
    if stop.status not in {"awaiting_resolution", "resolution_revision"}:
        raise HTTPException(status_code=409, detail="Для этой остановки сейчас не ожидается подтверждение устранения")
    count = _save_photos(db, stop, photos, "resolution")
    if count == 0:
        raise HTTPException(status_code=400, detail="Добавьте фото после устранения")
    stop.resolution_comment = _clean(comment, 5000)
    stop.resolution_submitted_at = utcnow()
    stop.status = "resolution_submitted"
    stop.updated_at = utcnow()
    _audit(db, stop, "worker", stop.initiator_name, "resolution_submitted", {"photos": count})
    db.commit()
    return _stop_out(stop, db, include_actions=True)


@router.get("/api/pb-mng/photos/{photo_id}")
def mobile_photo(photo_id: int, request: Request, db: Session = Depends(get_db)):
    photo = db.get(PbPhoto, photo_id)
    if not photo:
        raise HTTPException(status_code=404, detail="Фото не найдено")
    stop = db.get(PbStop, photo.stop_id)
    operator = _current_operator_or_none(request, db)
    if not operator:
        device_id, token_hash = _device_headers(request)
        if stop.device_id != device_id or stop.device_token_hash != token_hash:
            raise HTTPException(status_code=403, detail="Нет доступа к фото")
    path = PB_UPLOAD_DIR / photo.stored_name
    if not path.exists():
        raise HTTPException(status_code=404, detail="Файл фото не найден")
    return FileResponse(path, media_type=photo.mime_type, headers={"Cache-Control": "private, max-age=300"})


@router.post("/api/pb-mng/stops/{public_id}/pkm")
def submit_pkm(public_id: str, request: Request, text: str = Form(...), db: Session = Depends(get_db)):
    stop = _find_own_stop(db, public_id, request)
    if not stop.pkm_required or stop.status != "awaiting_pkm" or stop.pkm_status not in {"required", "revision"}:
        raise HTTPException(status_code=409, detail="ПКМ сейчас не ожидается")
    stop.pkm_text = _clean(text, 10000)
    if not stop.pkm_text:
        raise HTTPException(status_code=400, detail="Добавьте описание плана корректирующих мероприятий")
    stop.pkm_status = "submitted"
    stop.updated_at = utcnow()
    _audit(db, stop, "worker", stop.initiator_name, "pkm_submitted", {})
    db.commit()
    return _stop_out(stop, db, include_actions=True)


@router.get("/pb-mng/coordinator/", response_class=HTMLResponse, include_in_schema=False)
def coordinator_page(request: Request, db: Session = Depends(get_db)):
    operator = _current_operator_or_none(request, db)
    if not operator:
        return RedirectResponse("/login", status_code=303)
    return templates.TemplateResponse(request=request, name="pb_mng_coordinator.html", context={"operator": operator})


@router.get("/api/pb-mng/coordinator/stops")
def coordinator_stops(status: str = "", severity: str = "", q: str = "", limit: int = 200, db: Session = Depends(get_db), operator: Operator = Depends(coordinator_operator)):
    stmt = select(PbStop)
    if status:
        stmt = stmt.where(PbStop.status == status)
    if severity:
        stmt = stmt.where(PbStop.current_severity == severity)
    if q:
        token = f"%{q.strip()}%"
        stmt = stmt.where(or_(PbStop.public_id.ilike(token), PbStop.contractor.ilike(token), PbStop.location.ilike(token), PbStop.violation_text.ilike(token)))
    rows = db.scalars(stmt.order_by(PbStop.created_at.desc()).limit(min(max(limit, 1), 500))).all()
    return {"items": [_stop_out(x, db) for x in rows]}


@router.get("/api/pb-mng/coordinator/stops/{public_id}")
def coordinator_stop(public_id: str, db: Session = Depends(get_db), operator: Operator = Depends(coordinator_operator)):
    stop = db.scalar(select(PbStop).where(PbStop.public_id == public_id))
    if not stop:
        raise HTTPException(status_code=404, detail="Остановка не найдена")
    return _stop_out(stop, db, include_actions=True)


def _recommended_measures(severity: str) -> list[str]:
    if severity == "gross":
        return ["Работы остановлены до устранения", "Внесен в СТОП-ЛИСТ", "Направлен на ОБУЧЕНИЕ КБ"]
    if severity == "significant":
        return ["Работы остановлены до устранения", "Направлен на ОБУЧЕНИЕ КБ"]
    return ["Работы остановлены до устранения"]


@router.get("/api/pb-mng/coordinator/recommendations/{public_id}")
def recommendations(public_id: str, db: Session = Depends(get_db), operator: Operator = Depends(coordinator_operator)):
    stop = db.scalar(select(PbStop).where(PbStop.public_id == public_id))
    if not stop:
        raise HTTPException(status_code=404, detail="Остановка не найдена")
    severity = _effective_severity(stop, db)
    return {"severity": severity, "measures": _recommended_measures(severity), "suggested_block_days": 90 if severity == "gross" else 30 if severity == "significant" else None}


def _notification_recipients(db: Session, stop: PbStop) -> list[str]:
    ref = _reference_seed()
    result = []
    shared = _clean(ref.get("shared_stop_email"), 320)
    if shared:
        result.append(shared)
    routes = db.scalars(select(PbEmailRoute).where(PbEmailRoute.active.is_(True))).all()
    for route in routes:
        if route.block and route.block != stop.block:
            continue
        if route.contractor and route.contractor != stop.contractor:
            continue
        email = _clean(route.email, 320)
        if email and email not in result:
            result.append(email)
    return result


def _notify_verified(db: Session, stop: PbStop) -> None:
    recipients = _notification_recipients(db, stop)
    if not recipients:
        stop.mail_status = "no_recipients"
        return
    subject = f"Остановка работ {stop.public_id} — {SEVERITY_LABELS.get(stop.current_severity, stop.current_severity)}"
    body = "\n".join([
        f"Остановка: {stop.public_id}",
        f"Дата/время: {stop.occurred_at.astimezone().strftime('%d.%m.%Y %H:%M')}",
        f"Объект: {stop.field_name} / {stop.location}",
        f"Подрядчик: {stop.contractor}",
        f"Вид работ: {stop.work_type}",
        f"Нарушение: {stop.violation_text}",
        f"Категория: {SEVERITY_LABELS.get(stop.current_severity, stop.current_severity)}",
        f"Ответственный: {stop.responsible_fio} {stop.responsible_pass}",
        f"Меры: {', '.join(_json(stop.measures_json, []))}",
        "",
        f"Карточка координатора: https://rpo-mng.ru/pb-mng/coordinator/?stop={stop.public_id}",
    ])
    errors = []
    for recipient in recipients:
        try:
            send_notification(recipient, subject, body, idempotency_key=f"pb-stop-{stop.id}-{hashlib.sha1(recipient.encode()).hexdigest()[:12]}")
        except Exception as exc:
            errors.append(str(exc)[:160])
    stop.mail_status = "sent" if not errors else "partial_error:" + "; ".join(errors)[:500]


@router.post("/api/pb-mng/coordinator/stops/{public_id}/review")
def review_stop(public_id: str, payload: ReviewPayload, db: Session = Depends(get_db), operator: Operator = Depends(coordinator_operator)):
    stop = db.scalar(select(PbStop).where(PbStop.public_id == public_id))
    if not stop:
        raise HTTPException(status_code=404, detail="Остановка не найдена")
    if stop.status not in {"pending_verification", "returned_for_revision"}:
        raise HTTPException(status_code=409, detail="Карточка уже рассмотрена")

    action = payload.action.strip()
    if action == "return_for_revision":
        if not payload.note.strip():
            raise HTTPException(status_code=400, detail="Укажите, что необходимо доработать")
        stop.status = "returned_for_revision"
        stop.verification_note = _clean(payload.note, 5000)
        _audit(db, stop, "coordinator", operator.username, "returned_for_revision", {"note": stop.verification_note})
    elif action == "reject":
        if not payload.note.strip():
            raise HTTPException(status_code=400, detail="Укажите причину отклонения")
        stop.status = "rejected"
        stop.verification_note = _clean(payload.note, 5000)
        stop.closed_at = utcnow()
        _audit(db, stop, "coordinator", operator.username, "rejected", {"note": stop.verification_note})
    elif action == "verify":
        severity = payload.severity.strip() or _effective_severity(stop, db)
        if severity not in SEVERITY_LABELS:
            raise HTTPException(status_code=400, detail="Некорректная категория нарушения")
        stop.current_severity = severity
        stop.status = "awaiting_resolution"
        stop.verification_note = _clean(payload.note, 5000)
        selected_measures = payload.measures or _recommended_measures(severity)
        stop.measures_json = json.dumps(selected_measures, ensure_ascii=False)
        stop.course_name = _clean(payload.course_name, 300)
        stop.course_status = "required" if stop.course_name and "НЕ ТРЕБУЕТ" not in stop.course_name.upper() else "not_required"
        stop.pkm_required = bool(payload.pkm_required)
        stop.pkm_status = "required" if stop.pkm_required else "not_required"
        stop.verified_at = utcnow()
        stop.verified_by_id = operator.id
        should_block = bool(payload.block_responsible) or any("СТОП-ЛИСТ" in str(m).upper() for m in selected_measures)
        if should_block and stop.responsible_pass:
            ends_at = utcnow() + timedelta(days=payload.block_days) if payload.block_days else None
            stop.block_until = ends_at
            db.add(PbAccessRestriction(stop_id=stop.id, pass_number=normalize_pass_number(stop.responsible_pass) or stop.responsible_pass, fio=stop.responsible_fio, restriction_kind="personnel", reason=stop.violation_text, ends_at=ends_at, active=True))
        _audit(db, stop, "coordinator", operator.username, "verified", {"severity": severity, "measures": _json(stop.measures_json, []), "block": bool(should_block)})
        _notify_verified(db, stop)
    else:
        raise HTTPException(status_code=400, detail="Неизвестное решение координатора")

    stop.updated_at = utcnow()
    db.commit()
    return _stop_out(stop, db, include_actions=True)


@router.post("/api/pb-mng/coordinator/stops/{public_id}/resolution-review")
def review_resolution(public_id: str, payload: ResolutionReviewPayload, db: Session = Depends(get_db), operator: Operator = Depends(coordinator_operator)):
    stop = db.scalar(select(PbStop).where(PbStop.public_id == public_id))
    if not stop:
        raise HTTPException(status_code=404, detail="Остановка не найдена")
    if stop.status != "resolution_submitted":
        raise HTTPException(status_code=409, detail="Нет нового подтверждения устранения")
    if payload.action == "return":
        stop.status = "resolution_revision"
        stop.verification_note = _clean(payload.note, 5000)
        _audit(db, stop, "coordinator", operator.username, "resolution_returned", {"note": stop.verification_note})
    elif payload.action == "accept":
        stop.resolution_confirmed_at = utcnow()
        stop.status = _next_after_conditions(stop)
        if stop.status == "closed":
            stop.closed_at = utcnow()
        _audit(db, stop, "coordinator", operator.username, "resolution_accepted", {"next_status": stop.status})
    else:
        raise HTTPException(status_code=400, detail="Неизвестное решение")
    stop.updated_at = utcnow()
    db.commit()
    return _stop_out(stop, db, include_actions=True)


@router.post("/api/pb-mng/coordinator/stops/{public_id}/pkm")
def review_pkm(public_id: str, payload: PkmPayload, db: Session = Depends(get_db), operator: Operator = Depends(coordinator_operator)):
    stop = db.scalar(select(PbStop).where(PbStop.public_id == public_id))
    if not stop:
        raise HTTPException(status_code=404, detail="Остановка не найдена")
    if not stop.pkm_required:
        raise HTTPException(status_code=409, detail="ПКМ для остановки не требуется")
    stop.pkm_text = _clean(payload.text, 10000)
    if payload.action == "accept":
        stop.pkm_status = "accepted"
        if stop.resolution_confirmed_at:
            stop.status = _next_after_conditions(stop)
            if stop.status == "closed":
                stop.closed_at = utcnow()
        _audit(db, stop, "coordinator", operator.username, "pkm_accepted", {})
    elif payload.action == "return":
        stop.pkm_status = "revision"
        stop.status = "awaiting_pkm"
        _audit(db, stop, "coordinator", operator.username, "pkm_returned", {"note": stop.pkm_text})
    else:
        raise HTTPException(status_code=400, detail="Неизвестное решение по ПКМ")
    stop.updated_at = utcnow()
    db.commit()
    return _stop_out(stop, db, include_actions=True)


@router.post("/api/pb-mng/coordinator/stops/{public_id}/course-passed")
def course_passed(public_id: str, db: Session = Depends(get_db), operator: Operator = Depends(coordinator_operator)):
    stop = db.scalar(select(PbStop).where(PbStop.public_id == public_id))
    if not stop:
        raise HTTPException(status_code=404, detail="Остановка не найдена")
    if stop.course_status != "required":
        raise HTTPException(status_code=409, detail="Для остановки нет ожидающего курса")
    stop.course_status = "passed"
    if stop.resolution_confirmed_at and (not stop.pkm_required or stop.pkm_status == "accepted"):
        stop.status = _next_after_conditions(stop)
        if stop.status == "closed":
            stop.closed_at = utcnow()
    stop.updated_at = utcnow()
    _audit(db, stop, "coordinator", operator.username, "course_passed", {"course": stop.course_name, "next_status": stop.status})
    db.commit()
    return _stop_out(stop, db, include_actions=True)


@router.post("/api/pb-mng/coordinator/stops/{public_id}/unblock")
def unblock(public_id: str, db: Session = Depends(get_db), operator: Operator = Depends(coordinator_operator)):
    stop = db.scalar(select(PbStop).where(PbStop.public_id == public_id))
    if not stop:
        raise HTTPException(status_code=404, detail="Остановка не найдена")
    if not stop.resolution_confirmed_at:
        raise HTTPException(status_code=409, detail="Сначала подтвердите устранение нарушения")
    if stop.pkm_required and stop.pkm_status != "accepted":
        raise HTTPException(status_code=409, detail="Сначала подтвердите выполнение ПКМ")
    if stop.course_status == "required":
        raise HTTPException(status_code=409, detail="Сначала подтвердите прохождение назначенного курса")
    for restriction in stop.restrictions:
        if restriction.active:
            restriction.active = False
            restriction.released_at = utcnow()
            restriction.released_by_id = operator.id
    stop.status = "closed"
    stop.closed_at = utcnow()
    stop.updated_at = utcnow()
    _audit(db, stop, "coordinator", operator.username, "unblocked_and_closed", {})
    db.commit()
    return _stop_out(stop, db, include_actions=True)


@router.get("/api/pb-mng/coordinator/email-routes")
def list_email_routes(db: Session = Depends(get_db), operator: Operator = Depends(coordinator_operator)):
    rows = db.scalars(select(PbEmailRoute).order_by(PbEmailRoute.block, PbEmailRoute.contractor, PbEmailRoute.role)).all()
    return {"items": [{"id": x.id, "block": x.block, "contractor": x.contractor, "role": x.role, "recipient_name": x.recipient_name, "email": x.email, "active": x.active} for x in rows]}


@router.post("/api/pb-mng/coordinator/email-routes", status_code=201)
def add_email_route(payload: EmailRoutePayload, db: Session = Depends(get_db), operator: Operator = Depends(coordinator_operator)):
    email = payload.email.strip().lower()
    if "@" not in email or len(email) > 320:
        raise HTTPException(status_code=400, detail="Некорректный email")
    row = PbEmailRoute(block=_clean(payload.block, 160), contractor=_clean(payload.contractor, 300), role=_clean(payload.role, 80), recipient_name=_clean(payload.recipient_name, 240), email=email, active=True)
    db.add(row)
    db.commit()
    db.refresh(row)
    return {"id": row.id}


@router.get("/api/pb-mng/coordinator/stats")
def coordinator_stats(db: Session = Depends(get_db), operator: Operator = Depends(coordinator_operator)):
    total = db.scalar(select(func.count(PbStop.id))) or 0
    pending = db.scalar(select(func.count(PbStop.id)).where(PbStop.status == "pending_verification")) or 0
    active = db.scalar(select(func.count(PbStop.id)).where(PbStop.status.not_in(["closed", "rejected"]))) or 0
    closed = db.scalar(select(func.count(PbStop.id)).where(PbStop.status == "closed")) or 0
    by_severity = {s: db.scalar(select(func.count(PbStop.id)).where(PbStop.current_severity == s)) or 0 for s in SEVERITY_LABELS}
    return {"total": total, "pending": pending, "active": active, "closed": closed, "by_severity": by_severity}


def _excel_headers() -> list[str]:
    return [
        "№", "Дата", "Время", "Блок", "Подразделение", "Категория", "Структурное подразделение", "Месторождение", "Место производства работ", "Подрядная организация", "Субподрядная организация", "Вид работ", "№ наряда-допуска", "Барьер КБ", "Причина приостановки", "Ответственный / исполнители", "Должность / профессия", "№ пропуска ТС", "№ пропуска персонала", "Супервайзер / СК / строительный контроль", "Кто остановил", "Видеоаналитика", "Меры воздействия", "Приказ о дисциплинарном взыскании", "Отчет об устранении", "Дата возобновления", "Время возобновления", "Несанкционированное возобновление", "Назначенные курсы", "Курс пройден", "ПКМ", "Доступ персонала", "Доступ ТС", "Проводящий РПО"
    ]


@router.get("/api/pb-mng/coordinator/export.xlsx")
def export_stops(db: Session = Depends(get_db), operator: Operator = Depends(coordinator_operator)):
    rows = db.scalars(select(PbStop).order_by(PbStop.occurred_at.desc())).all()
    wb = Workbook()
    ws = wb.active
    ws.title = "СВОД"
    headers = _excel_headers()
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="0B3A73")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for stop in rows:
        _effective_severity(stop, db)
        restriction_active = any(r.active and (r.ends_at is None or (_as_utc(r.ends_at) or utcnow()) > utcnow()) for r in stop.restrictions)
        dt = stop.occurred_at.astimezone()
        closed = stop.closed_at.astimezone() if stop.closed_at else None
        ws.append([
            stop.public_id,
            dt.strftime("%d.%m.%Y"), dt.strftime("%H:%M"), stop.block, stop.initiator_unit,
            SEVERITY_LABELS.get(stop.current_severity, stop.current_severity), stop.structural_unit, stop.field_name,
            stop.location, stop.contractor, stop.subcontractor, stop.work_type, stop.permit_number, stop.violation_barrier,
            stop.violation_text, stop.responsible_fio, stop.responsible_position, "", stop.responsible_pass,
            stop.supervisor_fio, f"{stop.initiator_name} / {stop.initiator_unit}", "", ", ".join(_json(stop.measures_json, [])), "",
            "Да" if stop.resolution_confirmed_at else "Нет", closed.strftime("%d.%m.%Y") if closed else "", closed.strftime("%H:%M") if closed else "", "",
            stop.course_name, "Да" if stop.course_status in {"passed", "not_required"} else "Нет" if stop.course_status == "required" else "",
            "Да" if stop.pkm_status == "accepted" else "Не требуется" if stop.pkm_status == "not_required" else "Нет",
            "Запрещен" if restriction_active else "Разрешен", "", ""
        ])
    ws.freeze_panes = "A2"
    for col in range(1, len(headers) + 1):
        ws.column_dimensions[ws.cell(1, col).column_letter].width = 22
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    path = PB_EXPORT_DIR / f"PB_MNG_stops_{utcnow().strftime('%Y%m%d_%H%M%S')}.xlsx"
    wb.save(path)
    return FileResponse(path, filename=path.name, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


def active_pb_restrictions(db: Session, pass_number: str) -> list[dict]:
    canonical = normalize_pass_number(pass_number) or _clean(pass_number, 60)
    now = utcnow()
    rows = db.scalars(
        select(PbAccessRestriction, PbStop)
        .join(PbStop, PbStop.id == PbAccessRestriction.stop_id)
        .where(
            PbAccessRestriction.pass_number == canonical,
            PbAccessRestriction.active.is_(True),
            or_(PbAccessRestriction.ends_at.is_(None), PbAccessRestriction.ends_at > now),
        )
    ).all()
    # SQLAlchemy scalars() on a two-entity select returns only first entity, so fetch stop separately.
    result = []
    for restriction in rows:
        stop = db.get(PbStop, restriction.stop_id)
        result.append({
            "type": "pb_mng",
            "title": "Блокировка по остановке работ",
            "message": f"Остановка {stop.public_id if stop else restriction.stop_id}",
            "description": restriction.reason,
            "stop_id": stop.public_id if stop else "",
            "ends_at": restriction.ends_at.isoformat() if restriction.ends_at else None,
        })
    return result